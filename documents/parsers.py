"""
File content extraction utilities.

Each parser accepts raw file bytes and returns extracted text.
"""

import io
import logging

logger = logging.getLogger(__name__)


def parse_txt(file_bytes: bytes) -> str:
    """Extract text from a plain-text file."""
    return file_bytes.decode("utf-8", errors="replace")


def parse_pdf(file_bytes: bytes) -> str:
    """Extract text from all pages of a PDF using PyMuPDF."""
    import fitz  # PyMuPDF

    text_parts: list[str] = []
    with fitz.open(stream=file_bytes, filetype="pdf") as doc:
        for page in doc:
            text_parts.append(page.get_text())
    return "\n".join(text_parts)


def parse_docx(file_bytes: bytes) -> str:
    """Extract text from a Word .docx file (paragraphs + table cells)."""
    from docx import Document as DocxDocument

    doc = DocxDocument(io.BytesIO(file_bytes))

    parts: list[str] = []

    # Paragraphs
    for para in doc.paragraphs:
        if para.text.strip():
            parts.append(para.text)

    # Table cells
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    parts.append(cell.text)

    return "\n".join(parts)


def parse_xlsx(file_bytes: bytes) -> str:
    """
    Extract text from an Excel .xlsx file.
    Returns all cell values separated by tabs (columns) and newlines (rows),
    with sheet names as headers.
    """
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
    parts: list[str] = []

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        parts.append(f"=== {sheet_name} ===")
        for row in ws.iter_rows(values_only=True):
            cells = [str(c) if c is not None else "" for c in row]
            parts.append("\t".join(cells))

    wb.close()
    return "\n".join(parts)


def parse_file(file_bytes: bytes, file_type: int) -> str:
    """
    Dispatch to the right parser based on file_type enum value.
    file_type values: 1=TXT, 2=PDF, 3=DOCX, 4=XLSX
    """
    from .enums import FileType

    parsers = {
        FileType.TXT: parse_txt,
        FileType.PDF: parse_pdf,
        FileType.DOCX: parse_docx,
        FileType.XLSX: parse_xlsx,
    }

    parser = parsers.get(file_type)
    if parser is None:
        raise ValueError(f"Unsupported file type: {file_type}")

    return parser(file_bytes)
