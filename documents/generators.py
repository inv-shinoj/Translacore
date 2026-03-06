"""
Generate translated output files in the same format as the source.
Each generator accepts translated text and returns (bytes, content_type).
"""

import io
import os
from .enums import FileType


def generate_txt(translated_text: str, _filename: str) -> tuple[bytes, str]:
    """Generate a plain-text file from translated text."""
    return translated_text.encode("utf-8"), "text/plain; charset=utf-8"


def generate_pdf(translated_text: str, _filename: str) -> tuple[bytes, str]:
    """Generate a PDF file containing the translated text."""
    try:
        from reportlab.lib.pagesizes import A4
        from reportlab.pdfgen import canvas
        from reportlab.lib.units import inch
    except ImportError:
        # Fallback: just return text file if reportlab not installed
        return generate_txt(translated_text, _filename)

    buffer = io.BytesIO()
    c = canvas.Canvas(buffer, pagesize=A4)
    width, height = A4
    margin = inch
    y = height - margin
    line_height = 14

    # Simple text wrapping
    for line in translated_text.split("\n"):
        # Wrap long lines
        words = line.split()
        current_line = ""
        for word in words:
            test_line = f"{current_line} {word}".strip()
            if c.stringWidth(test_line, "Helvetica", 10) < (width - 2 * margin):
                current_line = test_line
            else:
                if current_line:
                    c.drawString(margin, y, current_line)
                    y -= line_height
                    if y < margin:
                        c.showPage()
                        y = height - margin
                current_line = word

        if current_line:
            c.drawString(margin, y, current_line)
            y -= line_height
            if y < margin:
                c.showPage()
                y = height - margin

    c.save()
    buffer.seek(0)
    return buffer.read(), "application/pdf"


def generate_docx(translated_text: str, _filename: str) -> tuple[bytes, str]:
    """Generate a Word .docx file from translated text."""
    from docx import Document

    doc = Document()
    for paragraph in translated_text.split("\n"):
        if paragraph.strip():
            doc.add_paragraph(paragraph)

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer.read(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def generate_xlsx(translated_text: str, _filename: str) -> tuple[bytes, str]:
    """
    Generate an Excel .xlsx file from translated text.
    Expects tab-separated columns and newline-separated rows,
    with optional "=== SheetName ===" headers.
    """
    from openpyxl import Workbook

    wb = Workbook()
    # Remove default sheet — we'll create named ones
    wb.remove(wb.active)

    current_sheet = None
    current_name = "Sheet1"

    for line in translated_text.split("\n"):
        # Check for sheet header
        if line.startswith("=== ") and line.endswith(" ==="):
            current_name = line[4:-4].strip()
            current_sheet = wb.create_sheet(title=current_name)
            continue

        if current_sheet is None:
            current_sheet = wb.create_sheet(title=current_name)

        cells = line.split("\t")
        current_sheet.append(cells)

    # If no sheets were created, create a default one
    if len(wb.sheetnames) == 0:
        ws = wb.create_sheet(title="Sheet1")
        for line in translated_text.split("\n"):
            ws.append([line])

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.read(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def generate_translated_file(
    translated_text: str,
    file_type: int,
    original_filename: str,
) -> tuple[bytes, str]:
    """
    Dispatch to the right generator based on file_type enum value.
    Returns (file_bytes, content_type).
    """
    generators = {
        FileType.TXT: generate_txt,
        FileType.PDF: generate_pdf,
        FileType.DOCX: generate_docx,
        FileType.XLSX: generate_xlsx,
    }

    generator = generators.get(file_type)
    if generator is None:
        # Fallback to plain text
        return generate_txt(translated_text, original_filename)

    return generator(translated_text, original_filename)
